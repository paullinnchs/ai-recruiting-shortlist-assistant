from pathlib import Path

SUPPORTED_RESUME_EXTENSIONS = {".txt", ".pdf", ".docx"}


def read_job_description(path: str = "input/job_description.txt") -> str:
    job_path = Path(path)
    if not job_path.exists():
        raise FileNotFoundError(f"Job description not found: {job_path}")

    text = job_path.read_text(encoding="utf-8").strip()
    if not text:
        raise ValueError("input/job_description.txt is empty.")
    return text


def read_txt_resumes(folder: str = "input/resumes") -> list[dict]:
    return read_resumes(folder)


def read_resumes(folder: str = "input/resumes") -> list[dict]:
    resumes_path = Path(folder)
    resumes_path.mkdir(parents=True, exist_ok=True)

    resumes = []
    for resume_file in sorted(resumes_path.iterdir()):
        if not resume_file.is_file():
            continue

        if resume_file.name.lower() == "sample_candidate.txt":
            continue

        if resume_file.suffix.lower() not in SUPPORTED_RESUME_EXTENSIONS:
            continue

        text = extract_resume_text(resume_file).strip()
        if text:
            resumes.append(
                {
                    "file_name": resume_file.name,
                    "candidate_name": guess_candidate_name(text, resume_file.stem),
                    "text": text,
                }
            )

    return resumes


def extract_resume_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension == ".txt":
        return path.read_text(encoding="utf-8")
    if extension == ".pdf":
        return extract_pdf_text(path)
    if extension == ".docx":
        return extract_docx_text(path)
    return ""


def extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    page_text = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            page_text.append(text)
    return "\n".join(page_text)


def extract_docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    table_cells = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    table_cells.append(cell.text)
    return "\n".join(paragraphs + table_cells)


def guess_candidate_name(text: str, fallback: str) -> str:
    for line in text.splitlines():
        cleaned = line.strip()
        if cleaned:
            return cleaned[:80]
    return fallback.replace("_", " ").replace("-", " ").title()
